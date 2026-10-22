"""Збирає звіт «Звіт_ЛР.docx» з лістингами реального коду та рамками під скріншоти."""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Звіт_ЛР.docx"

TEXT_FONT = "Times New Roman"
CODE_FONT = "Courier New"

doc = Document()

# ---------------------------------------------------------------- сторінка та стилі
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin = Cm(3), Cm(1.5)
    s.top_margin = s.bottom_margin = Cm(2)


def set_style_font(style, name, size, bold=None, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


set_style_font(doc.styles["Normal"], TEXT_FONT, 14)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15
set_style_font(doc.styles["Heading 1"], TEXT_FONT, 16, True, RGBColor(0, 0, 0))
set_style_font(doc.styles["Heading 2"], TEXT_FONT, 14, True, RGBColor(0, 0, 0))
set_style_font(doc.styles["Heading 3"], TEXT_FONT, 14, True, RGBColor(0, 0, 0))
set_style_font(doc.styles["List Bullet"], TEXT_FONT, 14)


# ---------------------------------------------------------------- допоміжні функції
def set_run_font(run, name, size=None):
    run.font.name = name
    if size:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def rich(par, text, size=None, italic=False, bold=False):
    """Текст у [[подвійних дужках]] підсвічується жовтим (це місця для заповнення)."""
    parts = text.split("[[")
    for i, part in enumerate(parts):
        if i == 0:
            chunks = [(part, False)]
        else:
            filled, _, rest = part.partition("]]")
            chunks = [(filled, True), (rest, False)]
        for chunk, highlight in chunks:
            if not chunk:
                continue
            run = par.add_run(chunk)
            set_run_font(run, TEXT_FONT, size)
            run.italic = italic
            run.bold = bold
            if highlight:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def p(text, align=None, bold=False, italic=False, size=None, indent=True):
    par = doc.add_paragraph()
    if indent:
        par.paragraph_format.first_line_indent = Cm(1.25)
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(par, text, size=size, italic=italic, bold=bold)
    return par


def bullets(items):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        rich(par, it)


def h(text, level=1):
    return doc.add_heading(text, level=level)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def code_lines(cell, lines, size=9):
    first = True
    for line in lines:
        par = cell.paragraphs[0] if first and not cell.paragraphs[0].text else cell.add_paragraph()
        first = False
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        run = par.add_run(line if line else " ")
        set_run_font(run, CODE_FONT, size)


def code(text, size=9):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, "F2F2F2")
    code_lines(cell, text.strip("\n").split("\n"), size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code_file(rel_path, size=8):
    full = ROOT / rel_path
    p(f"Файл {rel_path}:", indent=False, bold=True, size=12)
    code(full.read_text(encoding="utf-8").rstrip("\n"), size)


def caption(text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(par, text, size=12, italic=True)


def shot(n, title, commands, expect):
    """Жовта рамка під скріншот: клацніть усередину й вставте зображення (Ctrl+V)."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    row = t.rows[0]
    row.height = Cm(6.5)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    shade(cell, "FFF2CC")
    head = cell.paragraphs[0]
    run = head.add_run(f"▼ СКРІНШОТ №{n} — вставте зображення сюди й видаліть цей текст ▼")
    set_run_font(run, TEXT_FONT, 12)
    run.bold = True
    lab = cell.add_paragraph()
    r = lab.add_run("Команди для скріншота:")
    set_run_font(r, TEXT_FONT, 11)
    r.bold = True
    code_lines(cell, commands, 9)
    ex = cell.add_paragraph()
    r = ex.add_run("Має бути видно: " + expect)
    set_run_font(r, TEXT_FONT, 11)
    r.italic = True
    caption(f"Рисунок {n} — {title}")


def table(header, rows, widths=None, size=11):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, text in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, "D9D9D9")
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = c.paragraphs[0].add_run(text)
        set_run_font(r, TEXT_FONT, size)
        r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            rich(par, text, size=size)
    if widths:
        t.autofit = False
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def tcaption(text):
    par = doc.add_paragraph()
    par.paragraph_format.keep_with_next = True
    rich(par, text, size=12, italic=True)


# ================================================================ ТИТУЛЬНА СТОРІНКА
for line, bold in [
    ("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", True),
    ("[[Назва закладу освіти]]", False),
    ("[[Факультет / кафедра]]", False),
]:
    p(line, align="center", bold=bold, indent=False)
for _ in range(5):
    doc.add_paragraph()
p("ЗВІТ", align="center", bold=True, indent=False, size=20)
p("про виконання лабораторної роботи [[№]]", align="center", indent=False)
p("з дисципліни «[[назва дисципліни]]»", align="center", indent=False)
doc.add_paragraph()
p(
    "Тема: Розробка та тестування даних моделей з Django/Flask: "
    "від визначення до PyTest валідації",
    align="center", bold=True, indent=False,
)
for _ in range(5):
    doc.add_paragraph()
p("Виконав: студент групи [[група]]", indent=False)
p("[[Прізвище Ім'я По батькові]]", indent=False)
p("Перевірив: [[посада, прізвище викладача]]", indent=False)
for _ in range(4):
    doc.add_paragraph()
p("[[Місто]] — 2026", align="center", indent=False)
doc.add_page_break()

# ================================================================ МЕТА І СЕРЕДОВИЩЕ
h("Мета роботи")
p(
    "Набути практичних навичок проєктування моделей даних у Django ORM та Flask-SQLAlchemy, "
    "створення й застосування міграцій, налаштування зв'язків між сутностями, виконання "
    "CRUD-операцій і складних ORM-запитів, а також автоматизованого тестування за допомогою PyTest: "
    "фікстури, параметризація, мок-об'єкти та аналіз покриття коду.",
    align="justify",
)

h("Середовище виконання")
p("Для роботи створено два самостійні проєкти на спільній предметній області «Інтернет-бібліотека»:", align="justify")
bullets([
    "django_lib — проєкт Django (застосунок library), СУБД SQLite, файл db.sqlite3;",
    "flask_lib — проєкт Flask (Flask-SQLAlchemy + Flask-Migrate), СУБД SQLite, файл library.db.",
])
table(
    ["Компонент", "Версія"],
    [
        ["Операційна система", "[[вкажіть, напр. Windows 11]]"],
        ["Python", "[[python --version]]"],
        ["Django, pytest-django", "[[pip show django pytest-django]]"],
        ["Flask, Flask-SQLAlchemy, Flask-Migrate, pytest-flask", "[[pip show flask flask-sqlalchemy flask-migrate pytest-flask]]"],
        ["pytest, pytest-cov", "[[pip show pytest pytest-cov]]"],
    ],
    widths=[7.5, 9.5],
)
p("Встановлення залежностей (у віртуальному середовищі):", indent=False)
code(
    "cd django_lib\n"
    "python -m venv .venv\n"
    ".venv\\Scripts\\activate          # Linux/macOS:  source .venv/bin/activate\n"
    "pip install -r requirements.txt\n"
    "\n"
    "cd ../flask_lib\n"
    "pip install -r requirements.txt"
)

# ================================================================ ЗАВДАННЯ 1
doc.add_page_break()
h("Завдання 1. Проєктування та реалізація моделей Django")
h("1.1. Предметна область та сутності", 2)
p(
    "Обрано предметну область «Інтернет-бібліотека»: каталог книг, їхні автори та жанри, читачі "
    "з розширеним профілем і облік видачі книг. Проєкт створено з нуля, тому всі моделі є новими.",
    align="justify",
)
tcaption("Таблиця 1 — Сутності предметної області")
table(
    ["Сутність", "Призначення", "Основні поля"],
    [
        ["Author (Автор)", "Автор книги", "name, country (choices), birth_year, bio"],
        ["Genre (Жанр)", "Жанр, до якого належать книги", "name (unique), slug (unique), description"],
        ["Book (Книга)", "Позиція каталогу", "title, isbn (unique, 13 символів), author (FK), genres (M2M), year, copies_total, status (choices), created_at"],
        ["Reader (Читач)", "Зареєстрований користувач бібліотеки", "full_name, email (unique), registered_at"],
        ["ReaderProfile (Профіль читача)", "Додаткові дані читача", "reader (O2O), phone, address, birth_date, newsletter"],
        ["Loan (Видача)", "Факт видачі книги читачеві", "reader (FK), book (FK), loaned_at, due_date, returned_at"],
    ],
    widths=[4, 5, 8],
)

h("1.2. ER-діаграма (текстовий опис)", 2)
code(
    "+--------+ 1     N +--------+ N     M +--------+\n"
    "| Author |---------|  Book  |---------| Genre  |\n"
    "+--------+         +--------+         +--------+\n"
    "                       | 1\n"
    "                       |\n"
    "                       | N\n"
    "                   +--------+ N     1 +--------+ 1     1 +---------------+\n"
    "                   |  Loan  |---------| Reader |---------| ReaderProfile |\n"
    "                   +--------+         +--------+         +---------------+",
    size=10,
)
caption("Рисунок А — ER-діаграма предметної області (текстове подання)")

h("1.3. Зв'язки та політики on_delete", 2)
tcaption("Таблиця 2 — Зв'язки між сутностями")
table(
    ["Зв'язок", "Тип", "Реалізація у Django", "on_delete і обґрунтування"],
    [
        ["Author → Book", "один-до-багатьох", "Book.author = ForeignKey(Author, related_name='books')",
         "PROTECT: автора не можна видалити, поки в каталозі є його книги — інакше втрачається інформація про книги."],
        ["Book ↔ Genre", "багато-до-багатьох", "Book.genres = ManyToManyField(Genre, related_name='books')",
         "Django автоматично створює проміжну таблицю library_book_genres; одна книга має кілька жанрів, жанр містить багато книг."],
        ["Reader ↔ ReaderProfile", "один-до-одного", "ReaderProfile.reader = OneToOneField(Reader, related_name='profile')",
         "CASCADE: профіль не має сенсу без читача; UNIQUE на зовнішньому ключі гарантує «один-до-одного»."],
        ["Reader → Loan", "один-до-багатьох", "Loan.reader = ForeignKey(Reader, related_name='loans')",
         "CASCADE: при видаленні облікового запису читача видаляється й персональна історія його видач."],
        ["Book → Loan", "один-до-багатьох", "Loan.book = ForeignKey(Book, related_name='loans')",
         "PROTECT: книгу з історією видач видаляти не можна — потрібно архівувати (status='archived'). SET_NULL тут не підходить: видача без книги втрачає зміст."],
    ],
    widths=[3, 2.8, 5.2, 6],
    size=10,
)
p(
    "У кожної моделі визначено внутрішній клас Meta (ordering, verbose_name, verbose_name_plural) та метод "
    "__str__. Додатково: Book.get_absolute_url() пов'язує модель із маршрутом; Meta.indexes (індекс за роком "
    "видання) та UniqueConstraint для моделі Loan (унікальність комбінації reader + book + loaned_at).",
    align="justify",
)
h("1.4. Код моделей", 2)
p("Повний код моделей наведено нижче. Поле pages у Book закоментовано — його додають на кроці 2 (Завдання 2).", align="justify")
code_file("django_lib/library/models.py")

# ================================================================ ЗАВДАННЯ 2
doc.add_page_break()
h("Завдання 2. Міграції та управління схемою бази даних")
h("2.1. Початкова міграція", 2)
p(
    "Спочатку за моделями з Завдання 1 створюємо міграцію 0001_initial (makemigrations), "
    "потім застосовуємо її до порожньої бази (migrate).",
    align="justify",
)
code("cd django_lib\npython manage.py makemigrations library\npython manage.py migrate")
shot(
    1, "Створення та застосування початкової міграції",
    ["python manage.py makemigrations library", "python manage.py migrate"],
    "перелік Create model … та файл library/migrations/0001_initial.py; далі Applying library.0001_initial... OK.",
)

h("2.2. Друга міграція (зміна моделі)", 2)
p(
    "У моделі Book розкоментовано поле pages = PositiveIntegerField('Кількість сторінок', default=0). "
    "Поле має значення за замовчуванням, тому міграція застосовується до вже заповненої таблиці без помилок. "
    "Нова міграція створюється окремим файлом; уже застосовану 0001_initial вручну не змінюємо.",
    align="justify",
)
code(
    "python manage.py makemigrations library\n"
    "python manage.py migrate\n"
    "python manage.py showmigrations library"
)
shot(
    2, "Друга міграція та список міграцій (showmigrations)",
    ["python manage.py makemigrations library", "python manage.py migrate", "python manage.py showmigrations library"],
    "Add field pages to book, файл 0002_book_pages.py, Applying library.0002_book_pages... OK; у showmigrations обидві міграції зі значком [X].",
)

h("2.3. SQL міграції", 2)
code("python manage.py sqlmigrate library 0002")
shot(
    3, "SQL-код другої міграції (sqlmigrate)",
    ["python manage.py sqlmigrate library 0002"],
    "блок BEGIN; … COMMIT; з SQL-командами додавання стовпця pages (SQLite перебудовує таблицю library_book).",
)

h("2.4. Файли каталогу migrations/", 2)
tcaption("Таблиця 3 — Файли міграцій Django")
table(
    ["Файл", "Призначення"],
    [
        ["library/migrations/__init__.py", "Робить каталог Python-пакетом, щоб Django знаходив міграції."],
        ["library/migrations/0001_initial.py", "Створює таблиці Author, Genre, Book, Reader, ReaderProfile, Loan, проміжну таблицю book_genres, індекс book_year_idx і обмеження unique_loan_per_day."],
        ["library/migrations/0002_book_pages.py", "Операція AddField: додає у таблицю Book стовпець pages. Залежить від 0001_initial."],
        ["db.sqlite3 (таблиця django_migrations)", "Зберігає перелік застосованих міграцій, за ним Django визначає, що ще треба виконати."],
    ],
    widths=[6.5, 10.5],
)
h("2.5. Відтворюваність та відкат", 2)
p(
    "Ланцюжок відтворюється з нуля: достатньо видалити файл db.sqlite3 і виконати python manage.py migrate — "
    "Django послідовно застосує 0001 і 0002 на чистій базі. Відкат виконується командою "
    "python manage.py migrate library 0001 (після нього стовпець pages зникає; повторний migrate повертає його).",
    align="justify",
)

# ================================================================ ЗАВДАННЯ 3
doc.add_page_break()
h("Завдання 3. Моделі Flask-SQLAlchemy та Flask-Migrate")
h("3.1. Конфігурація та моделі", 2)
p(
    "СУБД — SQLite. URI підключення задається у config.py: «sqlite:///<абсолютний шлях до flask_lib>/library.db». "
    "Об'єкти db (SQLAlchemy) та migrate (Migrate) створені в app/extensions.py й підключаються у фабриці "
    "create_app(). Моделі узгоджені з Django-версією; відмінності: частину полів спрощено (немає status, created_at, birth_date) "
    "і використано явну проміжну таблицю book_genres, оголошену через db.Table.",
    align="justify",
)
tcaption("Таблиця 4 — Реалізація зв'язків у Flask-SQLAlchemy")
table(
    ["Зв'язок", "Реалізація"],
    [
        ["Author → Book (1:N)", "Book.author_id = db.ForeignKey('authors.id'); Author.books = db.relationship('Book', backref='author', lazy='select')"],
        ["Book ↔ Genre (M:N)", "Явна таблиця book_genres = db.Table(...) із двома зовнішніми ключами; Book.genres = db.relationship(..., secondary=book_genres, backref='books')"],
        ["Reader ↔ ReaderProfile (1:1)", "ReaderProfile.reader_id = db.ForeignKey(..., unique=True); Reader.profile = db.relationship(..., uselist=False, cascade='all, delete-orphan')"],
        ["Reader → Loan (1:N)", "Loan.reader_id = db.ForeignKey('readers.id'); Reader.loans = db.relationship('Loan', backref='reader', cascade='all, delete-orphan')"],
        ["Book → Loan (1:N)", "Loan.book_id = db.ForeignKey('books.id'); Loan.book = db.relationship('Book', backref='loans')"],
    ],
    widths=[4.5, 12.5],
    size=10,
)
code_file("flask_lib/app/models.py")

h("3.2. Початкова ревізія", 2)
code(
    "cd flask_lib\n"
    "flask --app run db init\n"
    'flask --app run db migrate -m "initial schema"\n'
    "flask --app run db upgrade"
)
shot(
    4, "Ініціалізація Flask-Migrate, перша ревізія та її застосування",
    ["flask --app run db init", 'flask --app run db migrate -m "initial schema"', "flask --app run db upgrade"],
    "Creating directory …migrations, Detected added table 'authors' … 'loans', Generating …_initial_schema.py, Running upgrade -> <id>, initial schema.",
)

h("3.3. Друга ревізія", 2)
p(
    "У моделі Book розкоментовано поле pages = db.Column(db.Integer, default=0). Генеруємо й застосовуємо другу ревізію, "
    "переглядаємо історію та структуру таблиць (show_schema.py читає схему через sqlalchemy.inspect).",
    align="justify",
)
code(
    'flask --app run db migrate -m "add book pages"\n'
    "flask --app run db upgrade\n"
    "flask --app run db history\n"
    "python show_schema.py"
)
shot(
    5, "Друга ревізія, історія ревізій і структура таблиць SQLite",
    ['flask --app run db migrate -m "add book pages"', "flask --app run db upgrade", "flask --app run db history", "python show_schema.py"],
    "Detected added column 'books.pages'; у history дві ревізії; у show_schema — таблиці authors, books (зі стовпцем pages), book_genres, genres, loans, readers, reader_profiles та їхні FK.",
)
h("3.4. Файли каталогу migrations/", 2)
table(
    ["Файл", "Призначення"],
    [
        ["migrations/alembic.ini, env.py, script.py.mako", "Службова конфігурація Alembic: підключення до бази з Flask-додатка та шаблон файлів ревізій."],
        ["migrations/versions/<id>_initial_schema.py", "Перша ревізія: create_table для всіх таблиць (upgrade) і drop_table (downgrade)."],
        ["migrations/versions/<id>_add_book_pages.py", "Друга ревізія: op.add_column('books', pages) та відкат op.drop_column; down_revision вказує на першу."],
        ["library.db (таблиця alembic_version)", "Зберігає ідентифікатор поточної ревізії."],
    ],
    widths=[6.5, 10.5],
)

h("3.5. Порівняння Flask-Migrate і Django Migrations", 2)
tcaption("Таблиця 5 — Порівняння процесів міграцій")
table(
    ["Аспект", "Django", "Flask-Migrate (Alembic)"],
    [
        ["Ініціалізація", "Вбудовано; потрібен лише застосунок у INSTALLED_APPS.", "Вручну один раз: flask db init."],
        ["Виявлення моделей", "Автоматично через INSTALLED_APPS.", "Моделі мають бути імпортовані (у фабриці створено import models)."],
        ["Генерація міграції", "makemigrations — автоматично.", "flask db migrate (autogenerate) — скрипт треба переглянути вручну."],
        ["Застосування", "migrate.", "flask db upgrade."],
        ["Відкат", "migrate <app> <номер>.", "flask db downgrade [-1 | ревізія]."],
        ["Перегляд SQL", "sqlmigrate.", "flask db upgrade --sql (офлайн-режим)."],
        ["Ідентифікатори", "Послідовні номери 0001, 0002…", "Хеші ревізій + посилання down_revision."],
        ["Перейменування поля", "Запитує, чи це rename.", "Бачить як drop + add: потрібно вручну замінити на op.alter_column / rename."],
        ["Облік застосованого", "Таблиця django_migrations.", "Таблиця alembic_version."],
        ["Злиття гілок", "Два «листки» міграцій — makemigrations --merge.", "Кілька heads — flask db merge heads -m '…'."],
    ],
    widths=[3.5, 6.5, 7],
    size=10,
)
p(
    "Під час злиття гілок обидві системи дають конфлікт, якщо дві гілки створили незалежні міграції від одного "
    "батьківського стану: у Django з'являються два «листки» в одному застосунку, у Alembic — кілька heads. "
    "Рішення — створити міграцію злиття, а також перевірити, що вона не суперечить змінам обох гілок; уже "
    "застосовані міграції не редагуємо. У Flask-Migrate додатково треба вручну перевіряти autogenerate-скрипт: "
    "він не помічає частину змін (перейменування, деякі зміни типів і обмежень).",
    align="justify",
)

# ================================================================ ЗАВДАННЯ 4
doc.add_page_break()
h("Завдання 4. ORM-запити: вибірки, агрегація, оптимізація")
h("4.1. Django: CRUD, Q-об'єкти, агрегація, select_related", 2)
p(
    "Запити зібрано в команді керування run_queries (код — у Додатку А). Вона спочатку очищає таблиці й заповнює їх "
    "тестовими даними, тому результат відтворюється. Запуск:",
    align="justify",
)
code("python manage.py run_queries")
shot(
    6, "Результат виконання python manage.py run_queries (Django ORM)",
    ["python manage.py run_queries"],
    "усі розділи CREATE / READ / AGGREGATE / ANNOTATE / N+1 / UPDATE / DELETE з результатами. Якщо вивід не вміщається, зробіть два знімки поспіль і вставте обидва під цим номером.",
)
tcaption("Таблиця 6 — Сценарії Django-запитів і очікувані результати (фактичний вивід — на Рис. 6)")
table(
    ["Операція", "Сценарій", "Очікуваний результат"],
    [
        ["CREATE", "create() для 3 авторів, 3 жанрів, 5 книг (+ genres.add), 3 читачів, 2 профілів, 4 видач", "Авторів: 3, жанрів: 3, книг: 5, читачів: 3, видач: 4"],
        ["filter()", "Книги з year < 1900", "«Захар Беркут» (1883), «Кобзар» (1840)"],
        ["exclude()", "Книги не українських авторів", "«1984», «Animal Farm»"],
        ["Q (OR)", "year ≥ 1940 АБО назва містить «Коб»", "«1984», «Animal Farm», «Кобзар»"],
        ["Q (NOT)", "Книги, що не належать до жанру «Поезія»", "«1984», «Animal Farm», «Захар Беркут»"],
        ["get()", "Книга за ISBN 9789660001003", "1984 (George Orwell)"],
        ["aggregate()", "Count, Avg(year), Sum(copies_total)", "{'books': 5, 'avg_year': 1904.4, 'copies': 9}"],
        ["annotate()", "Кількість книг у кожного автора", "George Orwell: 2, Іван Франко: 2, Тарас Шевченко: 1"],
        ["annotate() + filter", "Активні (не повернені) видачі на читача", "Олена Коваль: 2, Андрій Мельник: 1, Марія Шевчук: 0"],
        ["Зворотні зв'язки", "franko.books, dystopia.books, olena.loans, olena.profile", "Книги Франка: Захар Беркут, Мойсей; антиутопії: 1984, Animal Farm; у Олени 2 видачі, телефон +380501112233"],
        ["N+1", "Цикл по книгах з доступом до автора: без оптимізації, select_related, prefetch_related", "6 запитів (1 + 5) → 1 запит (JOIN) → 2 запити (prefetch жанрів)"],
        ["UPDATE", "save() для Кобзаря; update() з F('copies_total') + 1 для книг Orwell", "10 примірників; змінено рядків: 2; у «Animal Farm» — 3"],
        ["DELETE (PROTECT)", "Видалення автора, що має книги", "ProtectedError — автор залишається"],
        ["DELETE (умова)", "Видалення повернених видач", "(1, {'library.Loan': 1})"],
        ["DELETE (CASCADE)", "Видалення читача Андрія", "(3, …): читач, його профіль і його видача; лишилось читачів: 2, видач: 2"],
    ],
    widths=[3, 7, 7],
    size=10,
)

h("4.2. Flask: CRUD та транзакції через db.session", 2)
p(
    "Аналогічний набір операцій виконує скрипт demo_queries.py (код — у Додатку А). Додатково показано два "
    "транзакційні сценарії: успішну фіксацію кількох змін однією commit та скасування всіх змін через rollback "
    "при порушенні унікальності ISBN. Запуск (після міграцій):",
    align="justify",
)
code("python demo_queries.py")
shot(
    7, "Результат виконання demo_queries.py (SQLAlchemy ORM, commit/rollback)",
    ["python demo_queries.py"],
    "розділи CREATE / READ / AGGREGATE / GROUP BY / N+1 / UPDATE / ТРАНЗАКЦІЯ 1 / ТРАНЗАКЦІЯ 2 / DELETE; у транзакції 2 — повідомлення IntegrityError і «Тимчасовий читач у базі: 0».",
)
tcaption("Таблиця 7 — Сценарії Flask-запитів і очікувані результати (фактичний вивід — на Рис. 7)")
table(
    ["Операція", "Сценарій", "Очікуваний результат"],
    [
        ["CREATE", "add_all + flush() + commit() для тих самих даних", "Авторів: 3, жанрів: 3, книг: 5, читачів: 2, видач: 4"],
        ["filter()", "Book.year < 1900", "«Кобзар» (1840), «Захар Беркут» (1883)"],
        ["filter_by()", "Книги автора George Orwell (join Author)", "«1984», «Animal Farm»"],
        ["not_()", "Автор не з України", "«1984», «Animal Farm»"],
        ["or_()", "year ≥ 1940 АБО назва містить «Коб»", "«Кобзар», «Animal Farm», «1984» (за зростанням року)"],
        ["get / first", "db.session.get(Book, id); filter_by(isbn).first()", "Кобзар (Тарас Шевченко); 1984 (George Orwell)"],
        ["Агрегація", "func.count / avg / sum", "книг=5, середній рік=1904.4, примірників=9"],
        ["GROUP BY", "Кількість книг у автора (outerjoin + group_by)", "George Orwell: 2, Іван Франко: 2, Тарас Шевченко: 1"],
        ["Зв'язки", "franko.books, dystopia.books, olena.loans, olena.profile", "Книги Франка: Захар Беркут, Мойсей; антиутопії: 1984, Animal Farm; 2 видачі, телефон з профілю"],
        ["N+1", "Ледаче завантаження vs joinedload('author')", "Ледаче: 1 запит + по одному на кожного автора; joinedload: 1 запит"],
        ["UPDATE", "Зміна атрибута + commit; масове Query.update()", "10 примірників у «Кобзаря»; змінено рядків: 2"],
        ["Транзакція 1", "Читач + видача + зменшення copies_total, один commit", "Читач збережений: True, видач у Марії: 1"],
        ["Транзакція 2", "Додати читача й книгу з дубльованим ISBN, commit → помилка → rollback", "IntegrityError (UNIQUE …books.isbn); «Тимчасовий читач» у базі: 0"],
        ["DELETE", "Видалення автора з книгами; delete() за умовою; видалення читача (cascade)", "Автора видалити не вдалося (rollback); видалено повернених видач: 1; профілів: 1, видач: 3"],
    ],
    widths=[3, 7, 7],
    size=10,
)

h("4.3. Порівняння підходів", 2)
tcaption("Таблиця 8 — Django ORM і SQLAlchemy")
table(
    ["Аспект", "Django ORM", "Flask-SQLAlchemy"],
    [
        ["Збереження", "obj.save() / Model.objects.create() — запис одразу (autocommit).", "db.session.add() + db.session.commit() — патерн Unit of Work, явна фіксація."],
        ["Помилка в транзакції", "transaction.atomic().", "db.session.rollback() у блоці except."],
        ["Фільтрація", "filter(year__lt=1900), exclude(), Q(…) | Q(…), ~Q(…).", "filter(Book.year < 1900), or_(), not_(), filter_by()."],
        ["Агрегація", "aggregate(), annotate(Count('books')).", "func.count/avg/sum, group_by() та outerjoin вручну."],
        ["Зворотний зв'язок", "related_name на ForeignKey (author.books).", "backref у relationship."],
        ["Оптимізація N+1", "select_related (JOIN), prefetch_related (2 запити).", "joinedload / selectinload."],
        ["Порядок за замовчуванням", "Meta.ordering.", "Явний order_by() у кожному запиті."],
        ["Каскади", "Параметр on_delete на ключі.", "cascade='all, delete-orphan' на relationship + обмеження NOT NULL у базі."],
    ],
    widths=[3.5, 6.5, 7],
    size=10,
)

# ================================================================ ЗАВДАННЯ 5
doc.add_page_break()
h("Завдання 5. PyTest: тестові випадки, фікстури та параметризація")
h("5.1. Налаштування", 2)
p(
    "Для обох проєктів додано pytest.ini. У Django: DJANGO_SETTINGS_MODULE = config.settings (плагін pytest-django); "
    "у Flask — плагін pytest-flask, фікстура app у tests/conftest.py. Тести працюють на тестовій базі, що "
    "створюється й знищується автоматично: pytest-django створює окрему тестову базу, у Flask кожен тест "
    "отримує порожню in-memory SQLite (sqlite:///:memory:). Продуктивні бази db.sqlite3 та library.db не змінюються.",
    align="justify",
)
h("5.2. Структура каталогу tests/", 2)
code(
    "django_lib/                           flask_lib/\n"
    "├── pytest.ini                        ├── pytest.ini\n"
    "└── tests/                            └── tests/\n"
    "    ├── conftest.py   (фікстури)          ├── conftest.py   (фікстури)\n"
    "    ├── test_models.py                    ├── test_models.py\n"
    "    ├── test_views.py                     ├── test_routes.py\n"
    "    └── test_mocks.py                     └── test_mocks.py",
    size=9,
)
h("5.3. Фікстури та тести моделей", 2)
p(
    "Фікстури визначено в conftest.py: тестовий користувач (user), клієнт з авторизацією (auth_client), а також "
    "тестові записи пов'язаних моделей (author, genre, book, reader, profile, loan); у Flask — додаток app, "
    "обгортка db та ті самі записи. Параметризовані тести: test_loan_is_overdue (4 набори даних, включно з граничним "
    "випадком «термін спливає сьогодні») і test_isbn_length_validation у Django.",
    align="justify",
)
tcaption("Таблиця 9 — Тести моделей")
table(
    ["Тест", "Що перевіряє", "Django", "Flask"],
    [
        ["test_str_representations / test_str_and_repr", "Рядкове представлення", "так", "так"],
        ["test_book_default_values", "Значення за замовчуванням (copies_total=1, status)", "так", "так"],
        ["test_isbn_must_be_unique", "Обмеження unique → IntegrityError", "так", "так"],
        ["test_isbn_length_validation (parametrize)", "Валідація довжини ISBN (занадто короткий / довгий)", "так", "—"],
        ["test_many_to_many_books_and_genres", "Зв'язок «багато-до-багатьох» в обидва боки", "так", "так"],
        ["test_one_to_one_profile", "Зв'язок 1:1 та заборона другого профілю", "так", "так"],
        ["test_protect_… / test_author_with_books_cannot_be_deleted", "PROTECT / заборона видалення автора з книгами", "так", "так"],
        ["test_cascade_deletes_profile_and_loans", "Каскадне видалення профілю та видач", "так", "так"],
        ["test_loan_is_overdue (parametrize ×4)", "Логіка прострочення: прострочена / останній день / вчасно / повернена", "так", "так"],
    ],
    widths=[5.7, 6, 2.6, 2.7],
    size=10,
)
h("5.4. Запуск і результати", 2)
p("Команди запуску (разом із вимірюванням покриття для Завдання 7):", indent=False)
code(
    "cd django_lib\n"
    "pytest -v --cov=library --cov-report=term-missing\n"
    "\n"
    "cd ../flask_lib\n"
    "pytest -v --cov=app --cov-report=term-missing"
)
shot(
    8, "Django: вивід pytest -v та звіт покриття",
    ["cd django_lib", "pytest -v --cov=library --cov-report=term-missing"],
    "список тестів із PASSED, підсумок «N passed», нижче таблиця покриття (Name / Stmts / Miss / Cover / Missing). Якщо вивід задовгий — зробіть два знімки.",
)
shot(
    9, "Flask: вивід pytest -v та звіт покриття",
    ["cd flask_lib", "pytest -v --cov=app --cov-report=term-missing"],
    "список тестів із PASSED, підсумок «N passed», нижче таблиця покриття (Name / Stmts / Miss / Cover / Missing).",
)
p("Тести детерміновані: кожен працює на власній порожній базі, не залежить від порядку запуску й не використовує ручне заповнення даних.", align="justify")

# ================================================================ ЗАВДАННЯ 6
doc.add_page_break()
h("Завдання 6. Інтеграційне тестування views і маршрутів")
p(
    "HTTP-тести виконано тестовим клієнтом: у Django — фікстури client та auth_client (pytest-django), "
    "у Flask — client (pytest-flask) із передачею JSON через client.post(json=…). Тести охоплюють позитивні, "
    "негативні й граничні сценарії; їх результати видно на Рис. 8 і Рис. 9.",
    align="justify",
)
tcaption("Таблиця 10 — HTTP-тести Django")
table(
    ["Тест", "Маршрут", "Сценарій", "Очікуваний результат"],
    [
        ["test_book_list_ok_and_template", "GET /", "позитивний", "200, шаблон library/book_list.html, у тілі є «Кобзар»"],
        ["test_book_list_empty_catalog", "GET /", "граничний (порожній каталог)", "200, текст «Книг поки немає.»"],
        ["test_book_detail_ok", "GET /books/<pk>/", "позитивний", "200, у тілі автор і ISBN"],
        ["test_book_detail_not_found", "GET /books/9999/", "негативний", "404"],
        ["test_loans_redirects_anonymous_user", "GET /loans/", "негативний (не авторизований)", "302 на /admin/login/"],
        ["test_loans_available_for_authenticated_user", "GET /loans/", "позитивний (авторизований)", "200, у таблиці є «Олена Коваль»"],
    ],
    widths=[5, 3.2, 3.8, 5],
    size=10,
)
tcaption("Таблиця 11 — HTTP-тести Flask (REST API)")
table(
    ["Тест", "Маршрут", "Сценарій", "Очікуваний результат"],
    [
        ["test_list_books_ok", "GET /api/books", "позитивний", "200, JSON-список із «Кобзар»"],
        ["test_get_book_ok", "GET /api/books/<id>", "позитивний", "200, JSON з потрібним ISBN"],
        ["test_get_book_not_found", "GET /api/books/9999", "негативний", "404, {\"error\": \"Не знайдено\"}"],
        ["test_create_book_returns_201", "POST /api/books", "позитивний", "201, книгу збережено в базі"],
        ["test_create_book_missing_fields_returns_400 (×3)", "POST /api/books", "негативний (parametrize)", "400, «Відсутні поля»"],
        ["test_create_book_unknown_author_returns_400", "POST /api/books", "негативний", "400"],
        ["test_create_book_duplicate_isbn_returns_409", "POST /api/books", "граничний (дубль ISBN)", "409, rollback — у базі лишилась 1 книга"],
    ],
    widths=[5.2, 3.2, 3.6, 5],
    size=10,
)

# ================================================================ ЗАВДАННЯ 7
doc.add_page_break()
h("Завдання 7. Мок-тестування та аналіз покриття коду")
h("7.1. Мок-тести", 2)
tcaption("Таблиця 12 — Мок-тести та ізольовані залежності")
table(
    ["Тест (проєкт)", "Що замінено моком", "Тип", "Навіщо ізолювати", "Перевірка виклику"],
    [
        ["test_notify_overdue_sends_email (Django)", "library.services.send_mail", "зовнішня (пошта)", "Лист не надсилається реально; тест швидкий і не залежить від SMTP-сервера.", "assert_called_once_with(subject, message, from_email, recipient_list)"],
        ["test_is_overdue_with_mocked_today (Django)", "library.models.timezone.localdate", "внутрішня (час)", "Результат залежав би від дати запуску; мок «заморожує» сьогоднішню дату.", "assert_called_once_with()"],
        ["test_fetch_book_info_without_real_network (Flask)", "app.services.urlopen", "зовнішня (HTTP API Open Library)", "Мережа повільна, нестабільна й недоступна офлайн; відповідь підставляється вручну.", "assert_called_once_with(url, timeout=5)"],
        ["test_generate_card_number_with_mocked_random (Flask)", "app.services.random.randint", "внутрішня (випадковість)", "Випадкове число робить тест недетермінованим; мок повертає 123456.", "assert_called_once_with(100000, 999999)"],
    ],
    widths=[3.8, 3.2, 2.5, 4.3, 3.2],
    size=9,
)
p(
    "Моками підмінено лише зовнішні та побічні залежності (пошта, мережа, час, генератор випадкових чисел). "
    "Власну логіку, яку тестуємо (is_overdue, notify_overdue, fetch_book_info, generate_card_number), мок не замінює. "
    "Ефект: тести стають детермінованими, виконуються за мілісекунди й перевіряють саме логіку коду та правильність "
    "виклику залежності (аргументи).",
    align="justify",
)
h("7.2. Покриття коду", 2)
p(
    "Покриття виміряно командою pytest --cov (див. Рис. 8 та Рис. 9). Значення стовпців звіту coverage:",
    align="justify",
)
tcaption("Таблиця 13 — Колонки звіту coverage")
table(
    ["Колонка", "Значення"],
    [
        ["Name", "Модуль, для якого виконано вимірювання."],
        ["Stmts", "Кількість виконуваних операторів (statements) у модулі."],
        ["Miss", "Кількість операторів, що жодного разу не виконались під час тестів."],
        ["Cover", "Відсоток покриття: (Stmts − Miss) / Stmts × 100%."],
        ["Missing", "Номери рядків, що лишились без покриття (опція --cov-report=term-missing)."],
    ],
    widths=[3, 14],
)
h("7.3. Аналіз прогалин", 2)
p(
    "[[Звірте цей абзац зі своїм виводом coverage на Рис. 8 і Рис. 9 і за потреби скоригуйте.]] "
    "У Django найменше покрито management-команду library/management/commands/run_queries.py: це демонстраційний "
    "скрипт для Завдання 4, а не бізнес-логіка, тому відсутність для нього автотестів виправдана (його результат "
    "перевіряється візуально). Моделі, views, urls та services покриті майже повністю; міграції виконуються під час "
    "створення тестової бази. У Flask без покриття, ймовірно, лишились допоміжні методи app/models.py — __repr__ "
    "авторів, жанрів і читачів та __str__ профілю й видачі. Вони тривіальні, але за потреби закриваються простими "
    "тестами представлень. Файли demo_queries.py, show_schema.py і run.py лежать поза пакетом app й у вимірювання "
    "не входять. Непокриті гілки такого типу не впливають на працездатність API.",
    align="justify",
)

# ================================================================ ВИСНОВКИ
h("Висновки")
p(
    "У роботі створено два проєкти (Django та Flask) на спільній предметній області «Інтернет-бібліотека» з "
    "трьома типами зв'язків між сутностями, відтворюваними міграціями (по дві послідовні міграції в кожному "
    "проєкті), CRUD-операціями, агрегацією та оптимізацією запитів. Порівняння показало, що Django ORM надає "
    "більше «автоматики» (міграції, Meta.ordering, on_delete, select_related), тоді як SQLAlchemy вимагає явного "
    "керування сесією, транзакціями та запитами, але дає більше гнучкості. Автоматизоване тестування PyTest із "
    "фікстурами, параметризацією, HTTP-клієнтами й моками забезпечило детерміновану перевірку моделей та маршрутів, "
    "а звіт покриття показав, які частини коду ще не охоплені тестами.",
    align="justify",
)

# ================================================================ ДОДАТОК А
doc.add_page_break()
h("Додаток А. Лістинги коду")
p("Нижче наведено вихідний код ключових файлів обох проєктів.", indent=False)
h("А.1. Django", 2)
for f in [
    "django_lib/library/services.py",
    "django_lib/library/views.py",
    "django_lib/library/management/commands/run_queries.py",
    "django_lib/tests/conftest.py",
    "django_lib/tests/test_models.py",
    "django_lib/tests/test_views.py",
    "django_lib/tests/test_mocks.py",
]:
    code_file(f)
h("А.2. Flask", 2)
for f in [
    "flask_lib/app/__init__.py",
    "flask_lib/app/routes.py",
    "flask_lib/app/services.py",
    "flask_lib/demo_queries.py",
    "flask_lib/tests/conftest.py",
    "flask_lib/tests/test_models.py",
    "flask_lib/tests/test_routes.py",
    "flask_lib/tests/test_mocks.py",
]:
    code_file(f)

doc.save(OUT)
print("saved", OUT)
